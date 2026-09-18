"""
Green Time Ratio vs Deviation -- scatter plots, one per proxy pair
======================================================================
For each proxy pair (Flow_SMS, Flow_Snapshot, Flow_STOPLine, Occupancy),
produces one scatter-plot figure with THREE color- and marker-coded sets:
  - Mean Multiplier       (M1)  -> filled circle          'o'
  - 2nd-Degree Polynomial (M2)  -> filled triangle        '^'
  - GPR                   (M3)  -> circle with a cross    'o' + 'x' overlay

X axis : Green Time Ratio for that 30-second interval.
Y axis : Absolute deviation for that interval AFTER calibration --
         "MAD_After" for M1 and GPR, "MAD" for M2.

Run in PyCharm:
    pip install openpyxl numpy matplotlib
"""

import re
import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pathlib import Path

# ---------------------------------------------------------------------------
# CONFIG
# ---------------------------------------------------------------------------
density_file   = r"D:\Thesis_Final Data Analysis\05_12_25_1315_1345\30_sec\PCU\05_12_25_1315_1345_Combined_Density_30sec_T1_T2_Pooled_AllMethods.xlsx"
greentime_file = r"D:\Thesis_Final Data Analysis\05_12_25_1315_1345\GreenTime_GA_Dec05_1315.xlsx"

GREENTIME_SHEET = "Green Phases"
GREENTIME_HEADER_ROW = 4   # 0-indexed row containing the real column names

OUT_DIR = Path(density_file).parent   # PNGs saved alongside the density workbook

# One entry per proxy pair: (display name, M1 sheet, M2 sheet, GPR sheet)
PROXY_RUNS = [
    ("Flow_SMS",      "M. Mult Acc-Flow_SMS",      "2nd deg Acc-Flow_SMS",      "GPR Acc_Flow_SMS"),
    ("Flow_Snapshot", "M. Mult Acc-Flow_Snap",      "2nd deg Acc-Flow_Snap",      "GPR Acc_Flow_Snap"),
    ("Flow_STOPLine", "M. Mult Acc-Flow_STOPLine",  "2nd deg Acc-Flow_STOPLine",  "GPR Acc_Flow_STOPLine"),
    ("Occupancy",     "M. Mult Acc-Occ",            "2nd deg Acc-Occ",            "GPR Acc_Occ"),
]

METHOD_COLORS = {
    "Mean Multiplier":        "#1f77b4",   # blue
    "2nd-Degree Polynomial":  "#ff7f0e",   # orange
    "GPR":                    "#2ca02c",   # green
}

# Marker shape per method
METHOD_MARKERS = {
    "Mean Multiplier":        "o",   # circle
    "2nd-Degree Polynomial":  "^",   # triangle
    "GPR":                    "o",   # circle -- cross is drawn on top (see below)
}

POINT_SIZE = 35
POINT_ALPHA = 0.75
CROSS_COLOR = "#0b3d0b"        # dark green cross drawn inside the GPR circles
CROSS_SIZE_FACTOR = 0.45       # cross size relative to the circle area
CROSS_LINEWIDTH = 0.9

# ---------------------------------------------------------------------------
# Green-phase intervals -- loaded once, reused for every row/proxy/method
# ---------------------------------------------------------------------------
def load_green_phases(path, sheet=GREENTIME_SHEET, header_row=GREENTIME_HEADER_ROW):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    headers = [c.value for c in ws[header_row + 1]]
    col_idx = {h: i for i, h in enumerate(headers)}

    starts, ends = [], []
    for row in ws.iter_rows(min_row=header_row + 2, values_only=True):
        gs = row[col_idx["Green_Start_s"]]
        ge = row[col_idx["Green_End_s"]]
        if gs is None or ge is None:
            continue
        starts.append(float(gs))
        ends.append(float(ge))

    print(f"  {len(starts)} green phases loaded "
          f"(total green = {sum(e - s for s, e in zip(starts, ends)):.0f}s)")
    return np.array(starts), np.array(ends)


def green_ratio_for_window(t_start, t_end, green_starts, green_ends):
    """Fraction of [t_start, t_end] that overlaps any green phase."""
    overlap = np.maximum(
        0.0, np.minimum(t_end, green_ends) - np.maximum(t_start, green_starts)
    ).sum()
    return overlap / (t_end - t_start)


# ---------------------------------------------------------------------------
# Parse a row's time-range label into (t_start, t_end)
# ---------------------------------------------------------------------------
def parse_time_range(label):
    s = str(label).split(":")[-1].strip()
    a, b = s.split("-")
    return float(a), float(b)


# ---------------------------------------------------------------------------
# Read one calibration sheet -> list of (green_ratio, deviation) points.
# ---------------------------------------------------------------------------
def load_deviation_points(path, sheet_name, deviation_col_name,
                           green_starts, green_ends):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    headers = [c.value for c in ws[1]]
    col_idx = {h: i for i, h in enumerate(headers)}
    if deviation_col_name not in col_idx:
        raise KeyError(
            f"Column '{deviation_col_name}' not found in sheet '{sheet_name}'.\n"
            f"Available columns: {headers}"
        )
    dev_col = col_idx[deviation_col_name]

    green_ratios, deviations = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        time_s = row[1]
        ideal  = row[2]
        if time_s is None or ideal is None:
            continue
        label = row[0]
        dev = row[dev_col]
        if dev is None:
            continue
        t_start, t_end = parse_time_range(label)
        gr = green_ratio_for_window(t_start, t_end, green_starts, green_ends)
        green_ratios.append(gr)
        deviations.append(float(dev))

    return np.array(green_ratios), np.array(deviations)


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------
def main():
    print("Loading green-phase intervals ...")
    green_starts, green_ends = load_green_phases(greentime_file)

    for proxy_name, m1_sheet, m2_sheet, gpr_sheet in PROXY_RUNS:
        print(f"\n=== {proxy_name} ===")

        print(f"  Loading M1 (Mean Multiplier) -> '{m1_sheet}'")
        gr_m1, dev_m1 = load_deviation_points(
            density_file, m1_sheet, "MAD_After", green_starts, green_ends
        )

        print(f"  Loading M2 (2nd-Degree Polynomial) -> '{m2_sheet}'")
        gr_m2, dev_m2 = load_deviation_points(
            density_file, m2_sheet, "MAD", green_starts, green_ends
        )

        print(f"  Loading M3 (GPR) -> '{gpr_sheet}'")
        gr_gpr, dev_gpr = load_deviation_points(
            density_file, gpr_sheet, "MAD_After", green_starts, green_ends
        )

        fig, ax = plt.subplots(figsize=(9, 6.5))

        # --- M1: filled circles -------------------------------------------
        ax.scatter(gr_m1, dev_m1,
                   marker=METHOD_MARKERS["Mean Multiplier"],
                   color=METHOD_COLORS["Mean Multiplier"], s=POINT_SIZE,
                   alpha=POINT_ALPHA, edgecolors="white", linewidths=0.4,
                   label=f"Mean Multiplier (n={len(gr_m1)})", zorder=3)

        # --- M2: filled triangles -----------------------------------------
        ax.scatter(gr_m2, dev_m2,
                   marker=METHOD_MARKERS["2nd-Degree Polynomial"],
                   color=METHOD_COLORS["2nd-Degree Polynomial"],
                   s=POINT_SIZE * 1.15,     # triangles read smaller at equal area
                   alpha=POINT_ALPHA, edgecolors="white", linewidths=0.4,
                   label=f"2nd-Degree Polynomial (n={len(gr_m2)})", zorder=3)

        # --- M3: circle with a cross inside --------------------------------
        # Layer 1: the filled circle (this one carries the legend entry)
        gpr_circles = ax.scatter(
            gr_gpr, dev_gpr,
            marker="o",
            color=METHOD_COLORS["GPR"], s=POINT_SIZE,
            alpha=POINT_ALPHA, edgecolors="white", linewidths=0.4,
            zorder=3
        )
        # Layer 2: the cross drawn on top of each circle
        gpr_crosses = ax.scatter(
            gr_gpr, dev_gpr,
            marker="x",
            color=CROSS_COLOR, s=POINT_SIZE * CROSS_SIZE_FACTOR,
            linewidths=CROSS_LINEWIDTH, alpha=1.0,
            zorder=4
        )
        # Combine both layers into ONE legend entry so the key shows the
        # circle-with-cross symbol rather than two separate rows.
        gpr_handle = (gpr_circles, gpr_crosses)

        ax.set_xlabel("Green Time Ratio [-]", fontsize=11)
        ax.set_ylabel("Absolute Deviation from Ideal Density (PCU/km)", fontsize=11)
        ax.set_title(
            f"Green Time Ratio vs Deviation\n"
            f"Density (Accumulation) vs Density ({proxy_name})",
            fontsize=13, fontweight="bold"
        )
        ax.set_xlim(-0.03, 1.03)
        ax.set_ylim(bottom=0)
        ax.grid(True, linestyle=":", alpha=0.5)

        handles, labels = ax.get_legend_handles_labels()
        handles.append(gpr_handle)
        labels.append(f"GPR (n={len(gr_gpr)})")
        ax.legend(handles, labels,
                  loc="upper right", fontsize=9,
                  framealpha=0.9, edgecolor="#ccc",
                  scatterpoints=1,
                  handler_map={tuple: matplotlib.legend_handler.HandlerTuple(ndivide=None)})

        out_path = OUT_DIR / f"GreenRatio_vs_Deviation_{proxy_name}.png"
        fig.tight_layout()
        fig.savefig(out_path, dpi=200, bbox_inches="tight")
        plt.close(fig)
        print(f"  Saved -> {out_path}")

    print("\nDone. Four scatter plots were generated (one per proxy pair).")


if __name__ == "__main__":
    main()