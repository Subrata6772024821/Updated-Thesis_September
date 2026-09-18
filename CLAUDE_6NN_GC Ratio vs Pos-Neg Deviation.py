"""
Green Time Ratio vs SIGNED Deviation -- grouped unfilled bar charts, one per proxy pair
=======================================================================================
Reference  : Density_Accumulation (the "ideal").
Signed dev : Calibrated (Modified) Density  -  Density_Accumulation
             > 0  -> calibrated value OVERSHOOTS the ideal  ("positive deviation")
             < 0  -> calibrated value UNDERSHOOTS the ideal ("negative deviation")

For every Green Time Ratio bin, and separately for each method
(Mean Multiplier / 2nd-Degree Polynomial / GPR):
    - the POSITIVE signed deviations in that bin are averaged -> upward bar
    - the NEGATIVE signed deviations in that bin are averaged -> downward bar
Bars are drawn unfilled (coloured outline only), grouped side by side.

Four PNGs are produced, one per proxy pair.

Run in PyCharm:
    pip install openpyxl numpy matplotlib
"""

import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from pathlib import Path

# ---------------------------------------------------------------------------
# CONFIG
# ---------------------------------------------------------------------------
density_file   = r"D:\Thesis_Final Data Analysis\05_12_25_1315_1345\30_sec\PCU\05_12_25_1315_1345_Combined_Density_30sec_T1_T2_Pooled_AllMethods.xlsx"
greentime_file = r"D:\Thesis_Final Data Analysis\05_12_25_1315_1345\GreenTime_GA_Dec05_1315.xlsx"

GREENTIME_SHEET = "Green Phases"
GREENTIME_HEADER_ROW = 4          # 0-indexed row holding the real column names

OUT_DIR = Path(density_file).parent

PROXY_RUNS = [
    ("Flow_Snapshot", "M. Mult Acc-Flow_Snap",      "2nd deg Acc-Flow_Snap",      "GPR Acc_Flow_Snap"),
    ("Flow_SMS",      "M. Mult Acc-Flow_SMS",       "2nd deg Acc-Flow_SMS",       "GPR Acc_Flow_SMS"),
    ("Flow_STOPLine", "M. Mult Acc-Flow_STOPLine",  "2nd deg Acc-Flow_STOPLine",  "GPR Acc_Flow_STOPLine"),
    ("Occupancy",     "M. Mult Acc-Occ",            "2nd deg Acc-Occ",            "GPR Acc_Occ"),
]

METHOD_COLORS = {
    "Mean Multiplier":       "#1f77b4",
    "2nd-Degree Polynomial": "#ff7f0e",
    "GPR":                   "#2ca02c",
}
METHOD_ORDER = ["Mean Multiplier", "2nd-Degree Polynomial", "GPR"]

BIN_WIDTH = None          # None -> one group per distinct GT ratio; or e.g. 0.1
ROUND_DECIMALS = 3
BAR_EDGE_WIDTH = 1.4
GROUP_FILL = 0.82


def load_green_phases(path, sheet=GREENTIME_SHEET, header_row=GREENTIME_HEADER_ROW):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    headers = [c.value for c in ws[header_row + 1]]
    col_idx = {h: i for i, h in enumerate(headers)}
    starts, ends = [], []
    for row in ws.iter_rows(min_row=header_row + 2, values_only=True):
        gs, ge = row[col_idx["Green_Start_s"]], row[col_idx["Green_End_s"]]
        if gs is None or ge is None:
            continue
        starts.append(float(gs)); ends.append(float(ge))
    print(f"  {len(starts)} green phases loaded "
          f"(total green = {sum(e - s for s, e in zip(starts, ends)):.0f}s)")
    return np.array(starts), np.array(ends)


def green_ratio_for_window(t_start, t_end, gs, ge):
    overlap = np.maximum(0.0, np.minimum(t_end, ge) - np.maximum(t_start, gs)).sum()
    return overlap / (t_end - t_start)


def parse_time_range(label):
    s = str(label).split(":")[-1].strip()
    a, b = s.split("-")
    return float(a), float(b)


def find_col(headers, *substrings):
    for i, h in enumerate(headers):
        if h is None:
            continue
        hl = str(h).lower()
        if all(sub.lower() in hl for sub in substrings):
            return i
    raise KeyError(f"No column containing {substrings}. Headers: {headers[:12]}")


def load_signed_points(path, sheet_name, gs, ge):
    """-> (green_ratios, signed_deviation = calibrated - accumulation)"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows))
    acc_col = find_col(headers, "Density_Accumulation")
    cal_col = find_col(headers, "Modified Density")

    ratios, signed = [], []
    for row in rows:
        if row[1] is None or row[2] is None:
            continue
        acc, cal = row[acc_col], row[cal_col]
        if acc is None or cal is None:
            continue
        t0, t1 = parse_time_range(row[0])
        ratios.append(green_ratio_for_window(t0, t1, gs, ge))
        signed.append(float(cal) - float(acc))
    wb.close()
    return np.array(ratios), np.array(signed)


def bin_key(ratios):
    if BIN_WIDTH is None:
        return np.round(ratios, ROUND_DECIMALS)
    idx = np.floor(ratios / BIN_WIDTH).astype(int)
    idx = np.minimum(idx, int(round(1.0 / BIN_WIDTH)) - 1)
    return np.round((idx + 0.5) * BIN_WIDTH, ROUND_DECIMALS)


def main():
    print("Loading green-phase intervals ...")
    gs, ge = load_green_phases(greentime_file)

    for proxy_name, m1, m2, m3 in PROXY_RUNS:
        print(f"\n=== {proxy_name} ===")
        data = {}
        for label, sheet in zip(METHOD_ORDER, (m1, m2, m3)):
            r, d = load_signed_points(density_file, sheet, gs, ge)
            data[label] = (bin_key(r), d)
            print(f"  {label:<24} n={len(d):3d}  "
                  f"pos={np.sum(d > 0):3d}  neg={np.sum(d < 0):3d}")

        centers = np.unique(np.concatenate([k for k, _ in data.values()]))
        x = np.arange(len(centers))
        n = len(METHOD_ORDER)
        w = GROUP_FILL / n

        fig, ax = plt.subplots(figsize=(13, 6.8))
        for i, label in enumerate(METHOD_ORDER):
            keys, dev = data[label]
            offs = (i - (n - 1) / 2) * w
            pos_means, neg_means = [], []
            for c in centers:
                sel = dev[keys == c]
                p, q = sel[sel > 0], sel[sel < 0]
                pos_means.append(p.mean() if p.size else 0.0)
                neg_means.append(q.mean() if q.size else 0.0)
            col = METHOD_COLORS[label]
            ax.bar(x + offs, pos_means, width=w * 0.92, facecolor="none",
                   edgecolor=col, linewidth=BAR_EDGE_WIDTH, zorder=3)
            ax.bar(x + offs, neg_means, width=w * 0.92, facecolor="none",
                   edgecolor=col, linewidth=BAR_EDGE_WIDTH, zorder=3)

        ax.axhline(0, color="black", linewidth=1.0, zorder=4)
        ax.set_xticks(x)
        ax.set_xticklabels([f"{c:.2f}" for c in centers], rotation=90, fontsize=8)
        ax.set_xlabel("Green Time Ratio [-]", fontsize=11)
        ax.set_ylabel("Mean Signed Deviation from Ideal Density (PCU/km)", fontsize=11)
        ax.set_title(f"Green Time Ratio vs Signed Deviation\n"
                     f"Density (Accumulation) vs Density ({proxy_name})",
                     fontsize=13, fontweight="bold")
        ax.grid(True, axis="y", linestyle=":", alpha=0.5, zorder=0)
        ax.legend(handles=[Patch(facecolor="none", edgecolor=METHOD_COLORS[m],
                                 linewidth=BAR_EDGE_WIDTH, label=m)
                           for m in METHOD_ORDER],
                  loc="upper left", fontsize=9, framealpha=0.9, edgecolor="#ccc")
        ax.margins(x=0.01)

        out = OUT_DIR / f"GreenRatio_vs_SignedDeviation_{proxy_name}.png"
        fig.tight_layout()
        fig.savefig(out, dpi=200, bbox_inches="tight")
        plt.close(fig)
        print(f"  Saved -> {out}")

    print("\nDone.")


if __name__ == "__main__":
    main()